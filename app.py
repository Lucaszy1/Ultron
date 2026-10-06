"""Ultron: an AI-assisted bookkeeping helper for a QuickBooks-based business.

What this app actually does:
- Chat: forwards a conversation to Claude with a business-specific system prompt
  that is BUILT FROM a config file (so no company details are hardcoded).
- Spreadsheet tools: upload an Excel/CSV and get it parsed, column-mapped, and
  validated for sales orders, item codes, inventory, or duplicate detection
  (the real work lives in parsers.py).
- Export: write results back out to a formatted .xlsx.

The business context (company name, categories, code scheme, accounting rules)
is loaded from config.json, falling back to config.example.json. Copy the
example to config.json and edit it for your own business; config.json is
gitignored so nothing private is committed.
"""
import io
import json
import os

import openpyxl
import pandas as pd
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl.styles import Alignment, Font, PatternFill

import parsers

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

MODEL = os.environ.get("ULTRON_MODEL", "claude-sonnet-4-5")
MAX_TOKENS = int(os.environ.get("ULTRON_MAX_TOKENS", "2048"))
# Optional shared secret: if set, every /api-style POST must send it as
# X-App-Token. This stops a stranger who reaches the server from spending your
# Anthropic credits.
APP_TOKEN = os.environ.get("ULTRON_APP_TOKEN")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB upload cap

_client = None  # created lazily so importing this module (e.g. in tests) needs no API key


def get_client():
    global _client
    if _client is None:
        import anthropic
        _client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))
    return _client


# ---------- business config ----------
def load_business_config():
    for name in ("config.json", "config.example.json"):
        path = os.path.join(os.path.dirname(__file__), name)
        if os.path.exists(path):
            with open(path) as f:
                return json.load(f)
    return {}


CONFIG = load_business_config()


def build_system_prompt(guided=False, concise=False, industry=None):
    name = CONFIG.get("business_name", "the business")
    cats = ", ".join(CONFIG.get("inventory_categories", []))
    rules = "\n".join(f"- {r}" for r in CONFIG.get("accounting_rules", []))
    code_guide = CONFIG.get("code_guide", "")
    biz_industry = industry or CONFIG.get("industry", "")

    base = f"""You are Ultron, an AI QuickBooks Enterprise assistant for {name} ({biz_industry}).

You serve experienced QuickBooks users. Your job is to get things done: analyze
uploaded data, validate it, and return structured, actionable results.

Business context
- Inventory categories: {cats}
- Stock-code scheme: {code_guide}
- Accounting rules:
{rules}

When a spreadsheet is uploaded, parse it, map columns to QuickBooks fields,
validate required fields, and report exactly what is ready and what needs fixing.
For journal entries, give DEBIT/CREDIT lines with accounts, amounts, and a memo.
Be specific and actionable."""

    if guided:
        base += ("\n\nGUIDED MODE: ask one clarifying question at a time, confirm "
                 "before each step, and explain the reasoning as you go.")
    if concise:
        base += "\n\nBE CONCISE. Get to the answer immediately."
    return base


# ---------- auth guard ----------
@app.before_request
def check_token():
    if APP_TOKEN and request.method == "POST":
        if request.headers.get("X-App-Token") != APP_TOKEN:
            return jsonify({"error": "Unauthorized"}), 401


# ---------- pages ----------
@app.route("/")
def index():
    return render_template("index.html")


# ---------- chat ----------
@app.route("/config")
def business_config():
    """Expose non-sensitive business context so the UI can adapt (name,
    industry, category filters). Never returns secrets."""
    return jsonify({
        "business_name": CONFIG.get("business_name", ""),
        "industry": CONFIG.get("industry", ""),
        "inventory_categories": CONFIG.get("inventory_categories", []),
    })


@app.route("/chat", methods=["POST"])
def chat():
    data = request.get_json(silent=True) or {}
    messages = data.get("messages", [])
    system = build_system_prompt(
        guided=data.get("guided_mode", False),
        concise=data.get("response_length") == "concise",
        industry=data.get("industry"),
    )
    try:
        resp = get_client().messages.create(
            model=MODEL, max_tokens=MAX_TOKENS, system=system, messages=messages
        )
        return jsonify({"response": resp.content[0].text})
    except Exception as e:
        app.logger.exception("chat failed")
        return jsonify({"error": f"AI request failed: {e}"}), 502


# ---------- spreadsheet upload + parse ----------
PARSERS = {
    "sales_order": parsers.parse_sales_order,
    "stockcode": parsers.parse_item_codes,
    "item_codes": parsers.parse_item_codes,
    "inventory": parsers.parse_inventory,
    "duplicates": parsers.find_duplicates,
}


def _read_dataframe(file):
    """Read an uploaded .xlsx/.csv into a cleaned DataFrame."""
    if file.filename.lower().endswith(".csv"):
        content = file.read()
        try:
            df = pd.read_csv(io.BytesIO(content), header=0, encoding="utf-8", on_bad_lines="skip")
        except (UnicodeDecodeError, pd.errors.ParserError):
            df = pd.read_csv(io.BytesIO(content), header=0, encoding="latin-1", on_bad_lines="skip")
    else:
        df = pd.read_excel(file, header=0)
    return df.fillna("")


@app.route("/upload-excel", methods=["POST"])
def upload_excel():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    mode = request.form.get("mode", "sales_order")
    try:
        df = _read_dataframe(request.files["file"])
        columns = list(df.columns)
        rows = df.to_dict(orient="records")
        parser = PARSERS.get(mode)
        result = parser(columns, rows) if parser else {"columns": columns, "rows": rows[:50]}
        return jsonify({"success": True, "mode": mode, "total": len(rows), "result": result})
    except Exception as e:
        app.logger.exception("upload-excel failed")
        return jsonify({"error": f"Could not read file: {e}"}), 400


@app.route("/read-file-raw", methods=["POST"])
def read_file_raw():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    try:
        df = _read_dataframe(request.files["file"])
        text = "\n".join("\t".join(str(v) for v in row) for row in df.values)
        return jsonify({"text": text, "rows": len(df), "columns": list(df.columns)})
    except Exception as e:
        app.logger.exception("read-file-raw failed")
        return jsonify({"error": f"Could not read file: {e}"}), 400


# ---------- export ----------
@app.route("/export-excel", methods=["POST"])
def export_excel():
    data = request.get_json(silent=True) or {}
    rows = data.get("rows", [])
    columns = data.get("columns", [])
    filename = "".join(c for c in data.get("filename", "ultron_export") if c.isalnum() or c in "-_")
    title = (data.get("title", "Ultron Export") or "Export")[:31]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 2):
        for col_idx, col_name in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(width + 4, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"{filename or 'ultron_export'}.xlsx")


# ---------- AI analysis of parsed data ----------
ANALYSIS_PROMPTS = {
    "stockcode": "Analyze these item/stock codes. Identify: 1) codes that break the naming scheme, 2) missing fields, 3) duplicates, 4) new codes to add to QuickBooks.",
    "sales_order": "Analyze this parsed spreadsheet for QuickBooks sales-order import. Give: 1) what looks good, 2) exactly what needs fixing, 3) any stock codes that look wrong.",
    "item_codes": "Analyze these item codes. Identify: 1) codes that break the naming scheme, 2) missing fields, 3) duplicates, 4) new codes to add to QuickBooks.",
    "inventory": "Analyze this inventory. Give: 1) items critically low, 2) items at reorder point, 3) anomalies in quantity or cost, 4) recommended reorder quantities.",
    "duplicates": "Review these detected duplicates: 1) which are definitely duplicates, 2) which may be legitimate separate transactions, 3) recommended action for each.",
    "reconciliation": "Analyze this reconciliation between vendor data and QuickBooks inventory. For each discrepancy: 1) the exact issue, 2) the likely cause, 3) how to fix it in QuickBooks. Group by severity.",
}


@app.route("/analyze-upload", methods=["POST"])
def analyze_upload():
    data = request.get_json(silent=True) or {}
    mode = data.get("mode", "sales_order")
    parsed = data.get("parsed", {})
    instruction = ANALYSIS_PROMPTS.get(mode, ANALYSIS_PROMPTS["sales_order"])
    payload = json.dumps(parsed, indent=2)[:3000]
    try:
        resp = get_client().messages.create(
            model=MODEL, max_tokens=MAX_TOKENS,
            system=build_system_prompt(guided=data.get("guided_mode", False)),
            messages=[{"role": "user", "content": f"{instruction}\n\nData:\n{payload}"}],
        )
        return jsonify({"analysis": resp.content[0].text})
    except Exception as e:
        app.logger.exception("analyze-upload failed")
        return jsonify({"error": f"AI request failed: {e}"}), 502


if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "").lower() in ("1", "true", "yes")
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "5001"))
    print(f"\n[+] Ultron running at http://{host}:{port}\n")
    app.run(host=host, port=port, debug=debug)
